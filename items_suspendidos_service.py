import os
import re
import json
import unicodedata
from datetime import datetime
import openpyxl

FOLDER_ITEMS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Items suspendidos alumnos")
ALUMNOS_JSON = os.path.join(os.path.dirname(os.path.abspath(__file__)), "alumnos.json")
ITEMS_CORREGIDOS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "estado_items_corregidos.json")
EXPEDIENTES_REVISADOS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "estado_expedientes_revisados.json")
ITEMS_MANUALES_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "items_suspendidos_manuales.json")

def load_corregidos():
    if not os.path.exists(ITEMS_CORREGIDOS_FILE):
        return {}
    try:
        with open(ITEMS_CORREGIDOS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        print(f"Error cargando estado_items_corregidos.json: {e}")
        return {}

def save_corregidos(data):
    try:
        temp_file = ITEMS_CORREGIDOS_FILE + ".tmp"
        with open(temp_file, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        if os.path.exists(ITEMS_CORREGIDOS_FILE):
            os.replace(temp_file, ITEMS_CORREGIDOS_FILE)
        else:
            os.rename(temp_file, ITEMS_CORREGIDOS_FILE)
    except Exception as e:
        print(f"Error guardando estado_items_corregidos.json: {e}")

def load_expedientes_revisados():
    if not os.path.exists(EXPEDIENTES_REVISADOS_FILE):
        return {}
    try:
        with open(EXPEDIENTES_REVISADOS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        print(f"Error cargando estado_expedientes_revisados.json: {e}")
        return {}

def save_expedientes_revisados(data):
    try:
        temp_file = EXPEDIENTES_REVISADOS_FILE + ".tmp"
        with open(temp_file, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        if os.path.exists(EXPEDIENTES_REVISADOS_FILE):
            os.replace(temp_file, EXPEDIENTES_REVISADOS_FILE)
        else:
            os.rename(temp_file, EXPEDIENTES_REVISADOS_FILE)
    except Exception as e:
        print(f"Error guardando estado_expedientes_revisados.json: {e}")

def marcar_expediente_revisado(student_key, revisado=True):
    data = load_expedientes_revisados()
    k = normalize_str(student_key)
    if revisado:
        data[k] = {
            "fecha": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "student_key": student_key
        }
    else:
        if k in data:
            del data[k]
    save_expedientes_revisados(data)
    return data

def load_items_manuales():
    if not os.path.exists(ITEMS_MANUALES_FILE):
        return []
    try:
        with open(ITEMS_MANUALES_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        print(f"Error cargando items_suspendidos_manuales.json: {e}")
        return []

def save_items_manuales(data):
    try:
        temp_file = ITEMS_MANUALES_FILE + ".tmp"
        with open(temp_file, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        if os.path.exists(ITEMS_MANUALES_FILE):
            os.replace(temp_file, ITEMS_MANUALES_FILE)
        else:
            os.rename(temp_file, ITEMS_MANUALES_FILE)
    except Exception as e:
        print(f"Error guardando items_suspendidos_manuales.json: {e}")

def clean_lesson_str(lesson):
    if not lesson:
        return "GENERAL"
    s = normalize_str(lesson)
    s = re.sub(r'[^a-z0-9_-]', '_', s)
    s = re.sub(r'_+', '_', s).strip('_')
    return s[:40] or "GENERAL"

def agregar_item_manual(student_code, student_name, curso, item_name, nota=2.0, fecha=None, flight_id="", lesson="", observacion="", tipo=None):
    manuales = load_items_manuales()
    now_iso = datetime.now().strftime("%Y-%m-%d")
    fecha_use = parse_date(fecha) if fecha else now_iso
    s_key = student_code or normalize_str(student_name)
    clean_item = clean_item_text(item_name)
    norm_k = normalize_str(clean_item)
    clean_fl = clean_flight_str(flight_id) or "MANUAL"
    clean_les = clean_lesson_str(lesson) or "GENERAL"

    item_id = f"{s_key}__{curso}__{clean_fl}__{clean_les}__{norm_k}"

    nota_str = str(nota).strip()
    is_na = nota_str.upper() in ["N/A", "NA", "INC"]
    if tipo:
        tipo_final = tipo.upper()
    else:
        tipo_final = "ABIERTO" if is_na else "SUSPENDIDO"

    try:
        nota_val = "N/A" if is_na else float(nota_str)
    except ValueError:
        nota_val = "N/A"
        tipo_final = "ABIERTO"

    nuevo = {
        "id": item_id,
        "student_code": student_code,
        "student_name": student_name,
        "curso": curso or "General",
        "item": clean_item,
        "tipo": tipo_final,
        "nota": nota_val,
        "nota_display": str(nota_val),
        "fecha": fecha_use,
        "fecha_display": format_date_display(fecha_use),
        "flight_id": clean_fl,
        "lesson": lesson or "Añadido Manualmente",
        "recuperado": False,
        "con_vuelo_posterior": False,
        "detalle_recuperacion": observacion or "Registrado manualmente. Pendiente de reprogramación oficial.",
        "es_manual": True,
        "fecha_creacion": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    }

    # Reemplazar si ya existía el mismo ID o agregar
    manuales = [m for m in manuales if m.get("id") != item_id]
    manuales.append(nuevo)
    save_items_manuales(manuales)
    return nuevo

def eliminar_item_manual(item_id):
    manuales = load_items_manuales()
    manuales = [m for m in manuales if m.get("id") != item_id]
    save_items_manuales(manuales)
    return manuales

def clean_flight_str(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s

def check_if_resolved(corregidos_dict, item_id, student_key, course_name, flight_id, norm_key, fecha="", lesson=""):
    """
    Comprueba si un ítem está marcado como resuelto de forma persistente.
    Mantiene el estado marcado incluso cuando se importan nuevos Excels actualizados.
    Soporta formato con lección y formato legacy sin lección.
    """
    if not corregidos_dict:
        return False, ""
    
    # 1. Coincidencia exacta de ID
    if item_id in corregidos_dict:
        return True, corregidos_dict[item_id].get("fecha_marcado", "")
        
    # 2. Coincidencia con flight normalizado (sin .0)
    norm_flight = clean_flight_str(flight_id)
    alt_id = f"{student_key}__{course_name}__{norm_flight}__{norm_key}"
    if alt_id in corregidos_dict:
        return True, corregidos_dict[alt_id].get("fecha_marcado", "")

    # 2b. Coincidencia con lesson limpia en ID
    if lesson:
        les_clean = clean_lesson_str(lesson)
        alt_id_les = f"{student_key}__{course_name}__{norm_flight}__{les_clean}__{norm_key}"
        if alt_id_les in corregidos_dict:
            return True, corregidos_dict[alt_id_les].get("fecha_marcado", "")

    # 2c. Coincidencia flexible de prefijo/sufijo para claves con o sin lección
    prefix = f"{student_key}__{course_name}__{norm_flight}__"
    for cid in corregidos_dict:
        if cid.startswith(prefix) and (cid.endswith(f"__{norm_key}") or norm_key in cid):
            return True, corregidos_dict[cid].get("fecha_marcado", "")

    # 3. Coincidencia en metadata: student + curso + norm_item + fecha/flight
    for cid, entry in corregidos_dict.items():
        meta = entry.get("metadata", {})
        c_student = meta.get("student_key") or meta.get("student_code") or ""
        c_course = meta.get("curso") or ""
        c_item = meta.get("norm_key") or ""
        c_fecha = meta.get("fecha") or ""
        c_flight = clean_flight_str(meta.get("flight_id"))

        if c_student and normalize_str(c_student) == normalize_str(student_key):
            if c_course and c_course.upper() == course_name.upper():
                if not c_item or c_item.upper() == "ALL" or c_item == norm_key:
                    if (c_fecha and fecha and c_fecha == fecha) or (c_flight and norm_flight and c_flight == norm_flight):
                        return True, entry.get("fecha_marcado", "")
                    if not c_fecha and not c_flight and (c_item and c_item == norm_key):
                        return True, entry.get("fecha_marcado", "")

    return False, ""

def marcar_item_corregido(item_id, marcado=True, metadata=None):
    corregidos = load_corregidos()
    if marcado:
        parts = item_id.split("__")
        s_key = parts[0] if len(parts) > 0 else ""
        c_name = parts[1] if len(parts) > 1 else ""
        f_id = clean_flight_str(parts[2]) if len(parts) > 2 else ""
        if len(parts) >= 5:
            les_val = parts[3]
            n_key = parts[4]
        else:
            les_val = ""
            n_key = parts[3] if len(parts) > 3 else ""

        meta = dict(metadata or {})
        meta.setdefault("student_key", s_key)
        meta.setdefault("curso", c_name)
        meta.setdefault("flight_id", f_id)
        if les_val and "lesson" not in meta:
            meta["lesson"] = les_val
        meta.setdefault("norm_key", n_key)

        corregidos[item_id] = {
            "fecha_marcado": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "metadata": meta
        }
    else:
        if item_id in corregidos:
            del corregidos[item_id]
        # Limpieza de variantes con/sin .0
        norm_keys = [k for k in list(corregidos.keys()) if k.startswith(item_id) or item_id.startswith(k)]
        for nk in norm_keys:
            if nk in corregidos:
                del corregidos[nk]
    save_corregidos(corregidos)
    return corregidos

def marcar_lote_corregidos(item_ids, marcado=True):
    corregidos = load_corregidos()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for iid in item_ids:
        if marcado:
            parts = iid.split("__")
            s_key = parts[0] if len(parts) > 0 else ""
            c_name = parts[1] if len(parts) > 1 else ""
            f_id = clean_flight_str(parts[2]) if len(parts) > 2 else ""
            if len(parts) >= 5:
                les_val = parts[3]
                n_key = parts[4]
            else:
                les_val = ""
                n_key = parts[3] if len(parts) > 3 else ""

            meta = {
                "student_key": s_key,
                "curso": c_name,
                "flight_id": f_id,
                "lesson": les_val,
                "norm_key": n_key
            }
            corregidos[iid] = {
                "fecha_marcado": now_str,
                "metadata": meta
            }
        else:
            if iid in corregidos:
                del corregidos[iid]
    save_corregidos(corregidos)
    return corregidos

def normalize_str(s):
    if not s:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ASCII", "ignore").decode("ASCII")
    return s.strip().lower()

def clean_item_text(text):
    if not text:
        return ""
    s = str(text).strip()
    # Remove leading bullets, numbers or symbols like , •, -, etc.
    s = re.sub(r'^[•\-\*\—\?\s\d\.\)\(\/]+', '', s).strip()
    # Normalize multiple whitespace
    s = re.sub(r'\s+', ' ', s)
    return s

def load_alumnos_catalog():
    if not os.path.exists(ALUMNOS_JSON):
        return []
    try:
        with open(ALUMNOS_JSON, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        print(f"Error cargando alumnos.json: {e}")
        return []

def match_student_sheet(sheet_name, catalog):
    """Empareja el nombre de la pestaña (posiblemente truncado a 31 chars o con tildes distintas) con alumnos.json"""
    norm_sheet = normalize_str(sheet_name)
    if not norm_sheet:
        return {"code": "", "name": sheet_name}
    
    # 1. Exact match
    for a in catalog:
        if normalize_str(a.get("name")) == norm_sheet:
            return a
            
    # 2. Sheet name is prefix of catalog name (Excel 31 char truncation)
    for a in catalog:
        norm_a = normalize_str(a.get("name"))
        if norm_a.startswith(norm_sheet) or (len(norm_sheet) >= 15 and norm_sheet.startswith(norm_a)):
            return a

    # 3. Contains match
    for a in catalog:
        norm_a = normalize_str(a.get("name"))
        if norm_sheet in norm_a or norm_a in norm_sheet:
            return a

    # If no match in catalog, fallback to sheet name
    return {"code": "", "name": sheet_name}

def detect_course_from_workbook(wb, filename=""):
    """Detecta el curso (MEP, CPL, IR, PPL, etc.) inspeccionando el libro y las lecciones de sus hojas."""
    fname_upper = filename.upper()
    if "MEP" in fname_upper:
        return "MEP"
    if "CPL" in fname_upper:
        return "CPL"
    if "IR" in fname_upper:
        return "IR"

    # Search through sheets for lesson signatures
    for sname in wb.sheetnames[:5]:
        ws = wb[sname]
        for c in range(2, min(ws.max_column + 1, 30)):
            les = ws.cell(3, c).value
            if les:
                les_up = str(les).upper()
                if "ME." in les_up or "MEP" in les_up:
                    return "MEP"
                if "CPL" in les_up:
                    return "CPL"
                if "IR" in les_up or "ILS" in les_up or "NDB" in les_up:
                    return "IR"
                if "PPL" in les_up:
                    return "PPL"

    return "Curso General"

def parse_date(date_val):
    if isinstance(date_val, datetime):
        return date_val.strftime("%Y-%m-%d")
    if date_val:
        s = str(date_val).strip()
        # Look for YYYY-MM-DD
        m = re.search(r'(\d{4})[/-](\d{1,2})[/-](\d{1,2})', s)
        if m:
            return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
        # Look for DD/MM/YYYY
        m2 = re.search(r'(\d{1,2})[/-](\d{1,2})[/-](\d{4})', s)
        if m2:
            return f"{m2.group(3)}-{int(m2.group(2)):02d}-{int(m2.group(1)):02d}"
        return s[:10]
    return ""

def format_date_display(iso_date):
    if not iso_date:
        return ""
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', iso_date)
    if m:
        return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
    return iso_date

def format_lessons_summary(lessons_list):
    if not lessons_list:
        return ""
    clean_list = [l.strip() for l in lessons_list if l and l.strip()]
    if not clean_list:
        return ""
    if len(clean_list) == 1:
        return clean_list[0]
    
    short_codes = []
    for les in clean_list:
        m = re.match(r'^([A-Za-z0-9]+(?:\-[A-Za-z0-9]+)?(?:\.[0-9]+)*)', les)
        if m:
            code = m.group(1).rstrip('.')
            if code and code not in short_codes:
                short_codes.append(code)
        else:
            if les not in short_codes:
                short_codes.append(les)
    
    if len(short_codes) > 1:
        return " / ".join(short_codes)
    return " / ".join(clean_list)

def analyze_student_sheet(ws, student_info, course_name, corregidos_dict=None):
    """Analiza una hoja de alumno individual y detecta todos los suspensos y su estatus cronológico posterior."""
    if ws.max_row < 4 or ws.max_column < 2:
        return None

    if corregidos_dict is None:
        corregidos_dict = load_corregidos()

    # Step 1: Parse column headers (Row 1: Date, Row 2: Flight, Row 3: Lesson)
    columns_meta = []
    for col in range(2, ws.max_column + 1):
        raw_d = ws.cell(1, col).value
        iso_d = parse_date(raw_d)
        flight_id = ws.cell(2, col).value
        flight_str = str(flight_id).strip() if flight_id is not None else ""
        lesson_val = ws.cell(3, col).value
        lesson_str = str(lesson_val).strip() if lesson_val is not None else ""
        columns_meta.append({
            "col": col,
            "date": iso_d,
            "flight": flight_str,
            "lesson": lesson_str
        })

    # Step 2: Extract all evaluated items across rows
    # Map: normalized_item_name -> list of evaluations sorted by date/col
    items_evaluations = {}

    for r in range(4, ws.max_row + 1):
        raw_item = ws.cell(r, 1).value
        if not raw_item:
            continue
        clean_item = clean_item_text(raw_item)
        if not clean_item or len(clean_item) < 3:
            continue
        norm_key = normalize_str(clean_item)

        if norm_key not in items_evaluations:
            items_evaluations[norm_key] = {
                "display_name": clean_item,
                "evaluations": []
            }

        for meta in columns_meta:
            val = ws.cell(r, meta["col"]).value
            if val is None:
                continue
            raw_str = str(val).strip()
            # Try parsing numeric grade
            is_num = False
            num_grade = None
            try:
                num_grade = float(raw_str.replace(',', '.'))
                is_num = True
            except ValueError:
                pass

            if is_num:
                items_evaluations[norm_key]["evaluations"].append({
                    "col": meta["col"],
                    "date": meta["date"],
                    "flight": meta["flight"],
                    "lesson": meta["lesson"],
                    "type": "num",
                    "grade": num_grade,
                    "raw": raw_str
                })

    # Step 3: Identify failures (grade <= 2)
    failures_list = []

    for norm_key, item_data in items_evaluations.items():
        evals = item_data["evaluations"]
        if not evals:
            continue

        # Find evaluations requiring action: numeric <= 2.0 (suspendidos)
        target_evals = [e for e in evals if (e["type"] == "num" and e["grade"] <= 2.0)]
        if not target_evals:
            continue

        entries_processed = set()

        for f_idx, fail in enumerate(target_evals):
            fl_clean = clean_flight_str(fail["flight"])
            fl_key = fl_clean if fl_clean else clean_lesson_str(fail["lesson"])
            # Deduplicate by flight (or lesson if flight missing) to avoid repeating exercises across multi-lesson flights
            entry_key = (fail["date"], fl_key)
            if entry_key in entries_processed:
                continue
            entries_processed.add(entry_key)

            # Collect all distinct lessons for this flight on this date
            flight_lessons = []
            for e in evals:
                if e["date"] == fail["date"]:
                    efl = clean_flight_str(e["flight"]) if e["flight"] else clean_lesson_str(e["lesson"])
                    if efl == fl_key:
                        l_clean = (e["lesson"] or "").strip()
                        if l_clean and l_clean not in flight_lessons:
                            flight_lessons.append(l_clean)

            lesson_display = format_lessons_summary(flight_lessons) if flight_lessons else (fail["lesson"] or "")

            item_tipo = "SUSPENDIDO"

            # Look for subsequent evaluations on later dates or genuinely later flights
            later_evals = [
                e for e in evals 
                if (e["date"] > fail["date"] or (e["date"] == fail["date"] and clean_flight_str(e["flight"]) != fl_clean and e["col"] > fail["col"]))
            ]

            # Check if any later evaluation has numeric grade >= 3.0
            passed_subsequent = [e for e in later_evals if e["type"] == "num" and e["grade"] >= 3.0]
            
            if passed_subsequent:
                first_pass = passed_subsequent[0]
                con_vuelo_posterior = True
                fp_grade = int(first_pass['grade']) if first_pass['grade'].is_integer() else first_pass['grade']
                detalle = (
                    f"Comentario: Realizado en vuelo posterior #{first_pass['flight']} "
                    f"({first_pass['lesson'] or 'Misión posterior'}) con Nota {fp_grade} "
                    f"el {format_date_display(first_pass['date'])}. Requiere ser reprogramado y superado para darse por cerrado."
                )
                vuelo_post_info = {
                    "flight": first_pass["flight"],
                    "lesson": first_pass["lesson"],
                    "date": first_pass["date"],
                    "date_display": format_date_display(first_pass["date"]),
                    "grade": fp_grade
                }
            else:
                con_vuelo_posterior = False
                vuelo_post_info = None
                detalle = "Sin registro de vuelo posterior en este curso. Pendiente de reprogramar y superar."

            student_key = student_info.get('code') or normalize_str(student_info.get('name'))
            clean_fl = clean_flight_str(fail['flight'])
            item_id = f"{student_key}__{course_name}__{clean_fl or clean_lesson_str(lesson_display)}__{norm_key}"
            is_corregido, fecha_corregido = check_if_resolved(corregidos_dict, item_id, student_key, course_name, clean_fl, norm_key, fail['date'], lesson_display)

            failures_list.append({
                "id": item_id,
                "item": item_data["display_name"],
                "tipo": item_tipo,
                "nota": fail["grade"],
                "nota_display": str(fail["grade"]),
                "fecha": fail["date"],
                "fecha_display": format_date_display(fail["date"]),
                "flight_id": fail["flight"],
                "lesson": lesson_display,
                "curso": course_name,
                "recuperado": False,  # No se da por superado automáticamente por tener vuelo posterior
                "con_vuelo_posterior": con_vuelo_posterior,
                "vuelo_posterior_info": vuelo_post_info,
                "corregido_en_pr": is_corregido,
                "fecha_corregido": fecha_corregido,
                "detalle_recuperacion": detalle,
                "historial_intentos": [
                    {
                        "fecha": format_date_display(e["date"]),
                        "flight": e["flight"],
                        "lesson": e["lesson"],
                        "grade": e["grade"],
                        "type": e["type"]
                    } for e in evals
                ]
            })

    if not failures_list:
        return None

    # Count stats
    total_failures = len(failures_list)
    total_suspendidos_nota = sum(1 for f in failures_list if f["tipo"] == "SUSPENDIDO")
    total_abiertos = sum(1 for f in failures_list if f["tipo"] == "ABIERTO")
    total_con_vuelo_posterior = sum(1 for f in failures_list if f["con_vuelo_posterior"] and not f["corregido_en_pr"])
    corregidos_count = sum(1 for f in failures_list if f["corregido_en_pr"])
    activos_count = total_failures - corregidos_count

    return {
        "code": student_info.get("code") or "",
        "name": student_info.get("name") or ws.title,
        "sheet_name": ws.title,
        "curso": course_name,
        "total_suspensos": total_failures,
        "total_suspendidos_nota": total_suspendidos_nota,
        "total_items_abiertos": total_abiertos,
        "total_con_vuelo_posterior": total_con_vuelo_posterior,
        "suspensos_pendientes": activos_count,
        "suspensos_recuperados": total_con_vuelo_posterior,
        "total_corregidos_pr": corregidos_count,
        "total_activos_pr": activos_count,
        "items": failures_list
    }

def analyze_excel_file(filepath):
    """Analiza un único archivo Excel y devuelve los alumnos con suspensos."""
    if not os.path.exists(filepath):
        return {"error": "Archivo no encontrado", "alumnos": [], "curso": ""}

    wb = openpyxl.load_workbook(filepath, data_only=True)
    catalog = load_alumnos_catalog()
    corregidos_dict = load_corregidos()
    filename = os.path.basename(filepath)
    
    course_name = detect_course_from_workbook(wb, filename)

    students_with_failures = []

    for sname in wb.sheetnames:
        ws = wb[sname]
        student_info = match_student_sheet(sname, catalog)
        student_analysis = analyze_student_sheet(ws, student_info, course_name, corregidos_dict)
        if student_analysis:
            students_with_failures.append(student_analysis)

    # Sort students by name
    students_with_failures.sort(key=lambda s: normalize_str(s["name"]))

    total_susp = sum(s["total_suspensos"] for s in students_with_failures)
    total_corr = sum(s["total_corregidos_pr"] for s in students_with_failures)

    return {
        "filename": filename,
        "curso": course_name,
        "total_alumnos_con_suspensos": len(students_with_failures),
        "total_items_suspendidos": total_susp,
        "total_corregidos_pr": total_corr,
        "total_activos_pr": total_susp - total_corr,
        "alumnos": students_with_failures
    }

def analyze_all_courses(folder_path=FOLDER_ITEMS):
    """Analiza todos los archivos Excel en la carpeta y consolida la información por alumno."""
    if not os.path.exists(folder_path):
        os.makedirs(folder_path, exist_ok=True)
        return {"cursos": [], "alumnos": [], "resumen": {}}

    files = [f for f in os.listdir(folder_path) if f.endswith(".xlsx") and not f.startswith("~$")]
    if not files:
        return {"cursos": [], "alumnos": [], "resumen": {}}

    catalog = load_alumnos_catalog()
    corregidos_dict = load_corregidos()
    students_map = {} # key: student_code or normalized name
    courses_detected = set()
    files_analyzed = []

    for fname in files:
        fpath = os.path.join(folder_path, fname)
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
            course_name = detect_course_from_workbook(wb, fname)
            courses_detected.add(course_name)
            files_analyzed.append({"filename": fname, "curso": course_name})

            for sname in wb.sheetnames:
                ws = wb[sname]
                student_info = match_student_sheet(sname, catalog)
                res = analyze_student_sheet(ws, student_info, course_name, corregidos_dict)
                if not res:
                    continue

                student_key = student_info.get("code") or normalize_str(student_info.get("name"))
                if student_key not in students_map:
                    students_map[student_key] = {
                        "code": student_info.get("code") or "",
                        "name": student_info.get("name") or sname,
                        "cursos": set(),
                        "total_suspensos": 0,
                        "total_suspendidos_nota": 0,
                        "total_items_abiertos": 0,
                        "total_con_vuelo_posterior": 0,
                        "suspensos_pendientes": 0,
                        "suspensos_recuperados": 0,
                        "total_corregidos_pr": 0,
                        "total_activos_pr": 0,
                        "items": []
                    }

                students_map[student_key]["cursos"].add(course_name)
                students_map[student_key]["total_suspensos"] += res["total_suspensos"]
                students_map[student_key]["total_suspendidos_nota"] += res.get("total_suspendidos_nota", 0)
                students_map[student_key]["total_items_abiertos"] += res.get("total_items_abiertos", 0)
                students_map[student_key]["total_con_vuelo_posterior"] += res.get("total_con_vuelo_posterior", 0)
                students_map[student_key]["suspensos_pendientes"] += res["suspensos_pendientes"]
                students_map[student_key]["suspensos_recuperados"] += res.get("total_con_vuelo_posterior", 0)
                students_map[student_key]["total_corregidos_pr"] += res["total_corregidos_pr"]
                students_map[student_key]["total_activos_pr"] += res["total_activos_pr"]
                students_map[student_key]["items"].extend(res["items"])

        except Exception as e:
            print(f"Error procesando {fname}: {e}")

    # Incorporar ítems añadidos manualmente
    manuales = load_items_manuales()
    for m in manuales:
        s_code = m.get("student_code") or ""
        s_name = m.get("student_name") or ""
        s_key = s_code or normalize_str(s_name)
        c_name = m.get("curso") or "General"
        courses_detected.add(c_name)

        if s_key not in students_map:
            students_map[s_key] = {
                "code": s_code,
                "name": s_name,
                "cursos": set(),
                "total_suspensos": 0,
                "total_suspendidos_nota": 0,
                "total_items_abiertos": 0,
                "total_con_vuelo_posterior": 0,
                "suspensos_pendientes": 0,
                "suspensos_recuperados": 0,
                "total_corregidos_pr": 0,
                "total_activos_pr": 0,
                "items": []
            }

        students_map[s_key]["cursos"].add(c_name)
        is_corr, f_corr = check_if_resolved(corregidos_dict, m["id"], s_key, c_name, m.get("flight_id", "MANUAL"), normalize_str(m["item"]), m.get("fecha", ""), m.get("lesson", ""))
        
        m_copy = dict(m)
        m_copy["corregido_en_pr"] = is_corr
        m_copy["fecha_corregido"] = f_corr
        m_copy.setdefault("tipo", "ABIERTO" if str(m_copy.get("nota")).upper() in ["N/A", "NA"] else "SUSPENDIDO")
        m_copy.setdefault("nota_display", str(m_copy.get("nota")))

        # Check if student sheet has a matching item with subsequent flight evaluation
        if "con_vuelo_posterior" not in m_copy or not m_copy.get("con_vuelo_posterior"):
            norm_m_item = normalize_str(clean_item_text(m_copy.get("item", "")))
            for existing_item in students_map[s_key]["items"]:
                if normalize_str(clean_item_text(existing_item.get("item", ""))) == norm_m_item:
                    if existing_item.get("con_vuelo_posterior"):
                        m_copy["con_vuelo_posterior"] = True
                        m_copy["vuelo_posterior_info"] = existing_item.get("vuelo_posterior_info")
                        m_copy["detalle_recuperacion"] = existing_item.get("detalle_recuperacion")
                        break

        students_map[s_key]["total_suspensos"] += 1
        if m_copy["tipo"] == "ABIERTO":
            students_map[s_key]["total_items_abiertos"] += 1
        else:
            students_map[s_key]["total_suspendidos_nota"] += 1

        if m_copy.get("con_vuelo_posterior") and not is_corr:
            students_map[s_key]["total_con_vuelo_posterior"] += 1
            students_map[s_key]["suspensos_recuperados"] += 1

        if is_corr:
            students_map[s_key]["total_corregidos_pr"] += 1
        else:
            students_map[s_key]["total_activos_pr"] += 1
            students_map[s_key]["suspensos_pendientes"] += 1

        students_map[s_key]["items"].append(m_copy)

    # Cargar expedientes revisados
    revisados_dict = load_expedientes_revisados()

    # Convert student map to list
    final_students = []
    for k, s in students_map.items():
        s["cursos"] = sorted(list(s["cursos"]))
        # Sort items chronologically by date descending or by course
        s["items"].sort(key=lambda x: (x.get("curso", ""), x.get("fecha", "")), reverse=True)
        
        # Check si el expediente está revisado
        norm_k = normalize_str(k)
        norm_code = normalize_str(s.get("code", ""))
        norm_name = normalize_str(s.get("name", ""))
        
        is_rev = (norm_k in revisados_dict) or (norm_code and norm_code in revisados_dict) or (norm_name in revisados_dict)
        rev_entry = revisados_dict.get(norm_k) or revisados_dict.get(norm_code) or revisados_dict.get(norm_name) or {}
        
        s["expediente_revisado"] = is_rev
        s["fecha_expediente_revisado"] = rev_entry.get("fecha", "")

        final_students.append(s)

    final_students.sort(key=lambda s: normalize_str(s["name"]))

    total_suspensos = sum(s["total_suspensos"] for s in final_students)
    total_suspendidos_nota = sum(s.get("total_suspendidos_nota", 0) for s in final_students)
    total_items_abiertos = sum(s.get("total_items_abiertos", 0) for s in final_students)
    total_con_vuelo_posterior = sum(s.get("total_con_vuelo_posterior", 0) for s in final_students)
    total_pendientes = sum(s["suspensos_pendientes"] for s in final_students)
    total_corregidos = sum(s["total_corregidos_pr"] for s in final_students)
    total_activos = sum(s["total_activos_pr"] for s in final_students)
    total_revisados = sum(1 for s in final_students if s["expediente_revisado"])

    return {
        "cursos": sorted(list(courses_detected)),
        "archivos": files_analyzed,
        "total_alumnos_con_suspensos": len(final_students),
        "total_alumnos_revisados": total_revisados,
        "total_items_suspendidos": total_suspensos,
        "total_suspendidos_nota": total_suspendidos_nota,
        "total_items_abiertos": total_items_abiertos,
        "total_con_vuelo_posterior": total_con_vuelo_posterior,
        "total_recuperados": total_con_vuelo_posterior,
        "total_pendientes": total_pendientes,
        "total_corregidos_pr": total_corregidos,
        "total_activos_pr": total_activos,
        "alumnos": final_students
    }

def list_available_files(folder_path=FOLDER_ITEMS):
    if not os.path.exists(folder_path):
        os.makedirs(folder_path, exist_ok=True)
        return []
    files = [f for f in os.listdir(folder_path) if f.endswith(".xlsx") and not f.startswith("~$")]
    result = []
    for f in files:
        fpath = os.path.join(folder_path, f)
        stat = os.stat(fpath)
        # Try quick course detection
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
            curso = detect_course_from_workbook(wb, f)
            total_sheets = len(wb.sheetnames)
            wb.close()
        except:
            curso = "Desconocido"
            total_sheets = 0
            
        result.append({
            "filename": f,
            "curso": curso,
            "size": stat.st_size,
            "mtime": datetime.fromtimestamp(stat.st_mtime).strftime("%d/%m/%Y %H:%M"),
            "total_alumnos": total_sheets
        })
    return result
